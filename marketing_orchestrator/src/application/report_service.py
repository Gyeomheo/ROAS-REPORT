"""Marketing Orchestrator entrypoint."""

from __future__ import annotations

import hashlib
import os
import shutil
from datetime import date, datetime
from pathlib import Path
from time import perf_counter
from typing import Any, Dict, List

import polars as pl

from src.application.analysis_service import run_subsidiary_analysis
from src.application.reporting import lmdi as reporting_lmdi
from src.application.reporting import metrics as reporting_metrics
from src.application.reporting import rendering as reporting_rendering
from src.application.reporting import selectors as reporting_selectors
from src.infrastructure.excel_repository import (
    load_input_frame,
    save_html_calc_raw_workbook,
    save_output_workbook,
)
from src.infrastructure.report_exporter import save_summary_html, save_summary_json


DIMENSIONS: List[str] = ["SUBSIDIARY", "CHANNEL", "DIVISION", "PRODUCT"]
DEFAULT_CURRENT_YEAR = date.today().year
DEFAULT_PREVIOUS_YEAR = DEFAULT_CURRENT_YEAR - 1
CHANNEL_DIMS: List[str] = ["SUBSIDIARY", "CHANNEL"]
DIVISION_DIMS: List[str] = ["SUBSIDIARY", "CHANNEL", "DIVISION"]
MIN_TOPIC_ROAS_ABS_DELTA = 0.03
ScopeMaps = tuple[
    Dict[tuple[Any, ...], Dict[str, Any]],
    Dict[tuple[Any, ...], Dict[str, Any]],
    Dict[str, Dict[str, Any]],
]


def _env_int(name: str) -> int | None:
    raw = os.getenv(name)
    if raw is None or raw.strip() == "":
        return None
    try:
        return int(raw)
    except ValueError as exc:
        raise ValueError(f"Invalid integer for {name}: {raw}") from exc


def _resolve_comparison_years(curr_year: int | None, prev_year: int | None) -> tuple[int, int]:
    resolved_curr = curr_year if curr_year is not None else _env_int("ROAS_CURRENT_YEAR")
    resolved_prev = prev_year if prev_year is not None else _env_int("ROAS_PREVIOUS_YEAR")

    if resolved_curr is None and resolved_prev is None:
        resolved_curr = DEFAULT_CURRENT_YEAR
        resolved_prev = DEFAULT_PREVIOUS_YEAR
    elif resolved_curr is None:
        resolved_curr = int(resolved_prev) + 1
    elif resolved_prev is None:
        resolved_prev = int(resolved_curr) - 1

    if resolved_curr <= 0 or resolved_prev <= 0:
        raise ValueError(f"Comparison years must be positive: curr={resolved_curr}, prev={resolved_prev}")
    return int(resolved_curr), int(resolved_prev)


def _to_float(value: Any, default: float = 0.0) -> float:
    return reporting_metrics.to_float(value, default)


def _safe_ratio(num: float, den: float) -> float | None:
    return reporting_metrics.safe_ratio(num, den)


def _safe_pct_change(curr: float | None, prev: float | None) -> float | None:
    return reporting_metrics.safe_pct_change(curr, prev)


def _safe_ratio_expr(num: pl.Expr, den: pl.Expr) -> pl.Expr:
    return reporting_metrics.safe_ratio_expr(num, den)


def _safe_pct_change_expr(curr: pl.Expr, prev: pl.Expr) -> pl.Expr:
    return reporting_metrics.safe_pct_change_expr(curr, prev)


def _safe_log_ratio_expr(curr: pl.Expr, prev: pl.Expr) -> pl.Expr:
    return reporting_metrics.safe_log_ratio_expr(curr, prev)


def _log_mean_expr(curr: pl.Expr, prev: pl.Expr) -> pl.Expr:
    return reporting_metrics.log_mean_expr(curr, prev)


def _fmt_money(value: float | None) -> str:
    return reporting_metrics.fmt_money(value)


def _fmt_money_signed(value: float | None) -> str:
    return reporting_metrics.fmt_money_signed(value)


def _fmt_pct(value: float | None, signed: bool = True) -> str:
    return reporting_metrics.fmt_pct(value, signed=signed)


def _fmt_pct_yoy(value: float | None, signed: bool = True) -> str:
    return reporting_metrics.fmt_pct_yoy(value, signed=signed)


def _fmt_roas(value: float | None) -> str:
    return reporting_metrics.fmt_roas(value)


def _fmt_pp(value: float | None) -> str:
    return reporting_metrics.fmt_pp(value)


def _fmt_driver_metric(metric: str, value: float | None) -> str:
    return reporting_metrics.fmt_driver_metric(metric, value)


def _driver_label(driver: str | None) -> str:
    driver_text = str(driver or "").upper()
    if driver_text in {"CPC", "CVR", "AOV"}:
        return driver_text
    return "미확정"

def _trend(value: float | None, eps: float = 1e-9) -> str:
    return reporting_metrics.trend(value, eps=eps)


def _direct_contribution_pct(child_impact: float, parent_delta: float, mode_tag: str) -> float | None:
    return reporting_metrics.direct_contribution_pct(child_impact, parent_delta, mode_tag)

def _sum_aggregations() -> list[pl.Expr]:
    return reporting_lmdi.sum_aggregations()


def _sum_aggregations_full() -> list[pl.Expr]:
    return reporting_lmdi.sum_aggregations_full()


def _scope_maps(
    df: pl.DataFrame,
) -> tuple[Dict[tuple[Any, ...], Dict[str, Any]], Dict[tuple[Any, ...], Dict[str, Any]], Dict[str, Dict[str, Any]]]:
    return reporting_lmdi.scope_maps(df, channel_dims=CHANNEL_DIMS, division_dims=DIVISION_DIMS)


def _division_parallel_contribution(df: pl.DataFrame) -> List[Dict[str, Any]]:
    return reporting_lmdi.division_parallel_contribution(df)


def _final_drill_node(drill_trace: Any) -> Dict[str, Any]:
    if not isinstance(drill_trace, list) or not drill_trace:
        return {}
    node = drill_trace[-1]
    return node if isinstance(node, dict) else {}


def _trace_event_by_dim(drill_trace: Any, target_dim: str) -> Dict[str, Any]:
    if target_dim not in DIMENSIONS:
        return {}
    if not isinstance(drill_trace, list):
        return {}

    target_depth = DIMENSIONS.index(target_dim) + 1
    nodes: List[Dict[str, Any]] = []
    for item in drill_trace:
        if isinstance(item, dict):
            nodes.append(item)
        elif isinstance(item, list):
            for sub in item:
                if isinstance(sub, dict):
                    nodes.append(sub)

    for node in reversed(nodes):
        if not isinstance(node, dict):
            continue
        selected = node.get("selected_path", {})
        if not isinstance(selected, dict):
            continue
        depth = 0
        for dim in DIMENSIONS:
            if selected.get(dim) is None:
                break
            depth += 1
        if depth == target_depth and selected.get(target_dim) is not None:
            return node
    return {}


def _insight_text(
    row: Dict[str, Any],
    channel_map: Dict[tuple[Any, ...], Dict[str, Any]],
    division_map: Dict[tuple[Any, ...], Dict[str, Any]],
    subsidiary_map: Dict[str, Dict[str, Any]],
    impact_contribution_pct: float | None,
) -> tuple[str, str]:
    channel_key = tuple(row.get(dim) for dim in CHANNEL_DIMS)
    division_key = tuple(row.get(dim) for dim in DIVISION_DIMS)
    channel_row = channel_map.get(channel_key, {})
    division_row = division_map.get(division_key, {})

    channel_spend_curr = _to_float(channel_row.get("Spend_curr_sum"))
    channel_spend_prev = _to_float(channel_row.get("Spend_prev_sum"))
    channel_rev_curr = _to_float(channel_row.get("Revenue_curr_sum"))
    channel_rev_prev = _to_float(channel_row.get("Revenue_prev_sum"))
    channel_spend_yoy = _safe_pct_change(channel_spend_curr, channel_spend_prev)
    channel_rev_yoy = _safe_pct_change(channel_rev_curr, channel_rev_prev)
    channel_roas_curr = _safe_ratio(channel_rev_curr, channel_spend_curr)
    channel_roas_prev = _safe_ratio(channel_rev_prev, channel_spend_prev)
    channel_roas_yoy = _safe_pct_change(channel_roas_curr, channel_roas_prev)

    mode_tag = str(row.get("mode_tag", "") or "").upper()
    product_name = str(row.get("PRODUCT", "PRODUCT"))
    product_rev_curr = _to_float(row.get("Revenue_curr_sum"))
    product_rev_prev = _to_float(row.get("Revenue_prev_sum"))
    product_spend_curr = _to_float(row.get("Spend_curr_sum"))
    product_spend_prev = _to_float(row.get("Spend_prev_sum"))
    product_rev_yoy = _safe_pct_change(product_rev_curr, product_rev_prev)
    product_spend_yoy = _safe_pct_change(product_spend_curr, product_spend_prev)
    product_roas_curr = _safe_ratio(product_rev_curr, product_spend_curr)
    product_roas_prev = _safe_ratio(product_rev_prev, product_spend_prev)
    product_roas_yoy = _safe_pct_change(product_roas_curr, product_roas_prev)

    channel_name = str(row.get("CHANNEL", "CHANNEL"))
    headline = (
        f"{product_name} | 매출 {_fmt_money(product_rev_curr)}({_fmt_pct_yoy(product_rev_yoy)}) "
        f"| 투자 {_fmt_money(product_spend_curr)}({_fmt_pct_yoy(product_spend_yoy)}) "
        f"| ROAS {_fmt_roas(product_roas_curr)}({_fmt_pct_yoy(product_roas_yoy)})"
    )
    if mode_tag == "ISSUE":
        headline = f"{headline} [하락 유발]"
    elif mode_tag in {"IMPROVE", "DEFENSE"} and (channel_rev_yoy is not None and channel_rev_yoy < 0):
        headline = f"{headline} [채널 하락 상쇄]"

    division_spend_curr = _to_float(division_row.get("Spend_curr_sum"))
    division_spend_prev = _to_float(division_row.get("Spend_prev_sum"))
    division_rev_curr = _to_float(division_row.get("Revenue_curr_sum"))
    division_rev_prev = _to_float(division_row.get("Revenue_prev_sum"))
    division_clicks_curr = _to_float(division_row.get("Clicks_curr_sum"))
    division_clicks_prev = _to_float(division_row.get("Clicks_prev_sum"))
    division_orders_curr = _to_float(division_row.get("Orders_curr_sum"))
    division_orders_prev = _to_float(division_row.get("Orders_prev_sum"))
    division_spend_yoy = _safe_pct_change(division_spend_curr, division_spend_prev)
    division_rev_yoy = _safe_pct_change(division_rev_curr, division_rev_prev)
    division_roas_curr = _safe_ratio(division_rev_curr, division_spend_curr)
    division_roas_prev = _safe_ratio(division_rev_prev, division_spend_prev)
    division_roas_yoy = _safe_pct_change(division_roas_curr, division_roas_prev)
    division_cpc_curr = _safe_ratio(division_spend_curr, division_clicks_curr)
    division_cpc_prev = _safe_ratio(division_spend_prev, division_clicks_prev)
    division_cvr_curr = _safe_ratio(division_orders_curr, division_clicks_curr)
    division_cvr_prev = _safe_ratio(division_orders_prev, division_clicks_prev)
    division_aov_curr = _safe_ratio(division_rev_curr, division_orders_curr)
    division_aov_prev = _safe_ratio(division_rev_prev, division_orders_prev)
    division_rev_share = _safe_ratio(division_rev_curr, channel_rev_curr)
    channel_roas_delta = (
        0.0
        if channel_roas_curr is None or channel_roas_prev is None
        else _to_float(channel_roas_curr) - _to_float(channel_roas_prev)
    )
    channel_abs_delta = abs(channel_roas_delta)
    impact_adj = _to_float(row.get("Impact_adj"))
    division_name = str(row.get("DIVISION", "DIVISION"))
    subsidiary_name = str(row.get("SUBSIDIARY", ""))
    subsidiary_row = subsidiary_map.get(subsidiary_name, {})
    subsidiary_roas_curr = _safe_ratio(
        _to_float(subsidiary_row.get("Revenue_curr_sum")),
        _to_float(subsidiary_row.get("Spend_curr_sum")),
    )
    subsidiary_roas_prev = _safe_ratio(
        _to_float(subsidiary_row.get("Revenue_prev_sum")),
        _to_float(subsidiary_row.get("Spend_prev_sum")),
    )
    subsidiary_delta = (
        0.0
        if subsidiary_roas_curr is None or subsidiary_roas_prev is None
        else _to_float(subsidiary_roas_curr) - _to_float(subsidiary_roas_prev)
    )
    division_delta = (
        0.0
        if division_roas_curr is None or division_roas_prev is None
        else _to_float(division_roas_curr) - _to_float(division_roas_prev)
    )

    if impact_contribution_pct is None and channel_abs_delta > 0:
        impact_contribution_pct = _direct_contribution_pct(impact_adj, channel_roas_delta, mode_tag)

    if mode_tag == "ISSUE":
        contribution_rate_label = "상위 ROAS 하락 기여율"
    elif mode_tag in {"IMPROVE", "DEFENSE"}:
        contribution_rate_label = "상위 ROAS 개선 기여율"
    else:
        contribution_rate_label = "상위 ROAS 기여율"

    channel_event = _trace_event_by_dim(row.get("drill_trace"), "CHANNEL")
    division_event = _trace_event_by_dim(row.get("drill_trace"), "DIVISION")
    channel_impact = _to_float(channel_event.get("impact_adj"))
    division_impact = _to_float(division_event.get("impact_adj"))
    channel_direct_pct = _direct_contribution_pct(channel_impact, subsidiary_delta, mode_tag)
    division_direct_pct = _direct_contribution_pct(division_impact, channel_roas_delta, mode_tag)
    product_direct_pct = _direct_contribution_pct(impact_adj, division_delta, mode_tag)
    if product_direct_pct is None:
        product_direct_pct = impact_contribution_pct

    product_clicks_curr = _to_float(row.get("Clicks_curr_sum"))
    product_clicks_prev = _to_float(row.get("Clicks_prev_sum"))
    product_orders_curr = _to_float(row.get("Orders_curr_sum"))
    product_orders_prev = _to_float(row.get("Orders_prev_sum"))

    product_cpc_curr = row.get("CPC_curr")
    product_cpc_prev = row.get("CPC_prev")
    product_cvr_curr = row.get("CVR_curr")
    product_cvr_prev = row.get("CVR_prev")
    product_aov_curr = row.get("AOV_curr")
    product_aov_prev = row.get("AOV_prev")
    if product_cpc_curr is None:
        product_cpc_curr = _safe_ratio(product_spend_curr, product_clicks_curr)
    if product_cpc_prev is None:
        product_cpc_prev = _safe_ratio(product_spend_prev, product_clicks_prev)
    if product_cvr_curr is None:
        product_cvr_curr = _safe_ratio(product_orders_curr, product_clicks_curr)
    if product_cvr_prev is None:
        product_cvr_prev = _safe_ratio(product_orders_prev, product_clicks_prev)
    if product_aov_curr is None:
        product_aov_curr = _safe_ratio(product_rev_curr, product_orders_curr)
    if product_aov_prev is None:
        product_aov_prev = _safe_ratio(product_rev_prev, product_orders_prev)

    driver = str(row.get("primary_driver", "NONE") or "NONE").upper()
    if driver == "AOV":
        product_driver_curr = product_aov_curr
        product_driver_prev = product_aov_prev
        division_driver_curr = division_aov_curr
        division_driver_prev = division_aov_prev
        driver_dlog = row.get("dlog_AOV")
    elif driver == "CVR":
        product_driver_curr = product_cvr_curr
        product_driver_prev = product_cvr_prev
        division_driver_curr = division_cvr_curr
        division_driver_prev = division_cvr_prev
        driver_dlog = row.get("dlog_CVR")
    elif driver == "CPC":
        product_driver_curr = product_cpc_curr
        product_driver_prev = product_cpc_prev
        division_driver_curr = division_cpc_curr
        division_driver_prev = division_cpc_prev
        driver_dlog = row.get("dlog_CPC")
    else:
        driver = "NONE"
        product_driver_curr = None
        product_driver_prev = None
        division_driver_curr = None
        division_driver_prev = None
        driver_dlog = None

    driver_yoy = _safe_pct_change(product_driver_curr, product_driver_prev)
    division_driver_yoy = _safe_pct_change(division_driver_curr, division_driver_prev)
    driver_gap_pp = None
    if driver_yoy is not None and division_driver_yoy is not None:
        driver_gap_pp = driver_yoy - division_driver_yoy

    dlog_cvr_abs = abs(_to_float(row.get("dlog_CVR")))
    dlog_aov_abs = abs(_to_float(row.get("dlog_AOV")))
    dlog_cpc_abs = abs(_to_float(row.get("dlog_CPC")))
    driver_dlog_abs = abs(_to_float(driver_dlog))
    signal_max_abs = max(dlog_cvr_abs, dlog_aov_abs, dlog_cpc_abs)
    tag = str(row.get("tag", "") or "").upper()
    driver_label = _driver_label(driver)

    channel_text = (
        f"1) 채널 {channel_name}: 매출 {_fmt_money(channel_rev_curr)}({_fmt_pct_yoy(channel_rev_yoy)}), "
        f"투자 {_fmt_money(channel_spend_curr)}({_fmt_pct_yoy(channel_spend_yoy)}), "
        f"ROAS {_fmt_roas(channel_roas_curr)}({_fmt_pct_yoy(channel_roas_yoy)}), "
        f"{contribution_rate_label} {_fmt_pct(channel_direct_pct, signed=False)}"
    )
    division_text = (
        f"2) BU {division_name}: 채널 내 매출 비중 {_fmt_pct(division_rev_share, signed=False)}, "
        f"매출 {_fmt_money(division_rev_curr)}({_fmt_pct_yoy(division_rev_yoy)}), "
        f"투자 {_fmt_money(division_spend_curr)}({_fmt_pct_yoy(division_spend_yoy)}), "
        f"ROAS {_fmt_roas(division_roas_curr)}({_fmt_pct_yoy(division_roas_yoy)}), "
        f"{contribution_rate_label} {_fmt_pct(division_direct_pct, signed=False)}"
    )
    product_text = (
        f"3) 제품 {product_name}: {contribution_rate_label} {_fmt_pct(product_direct_pct, signed=False)}, "
        f"핵심 지표 {driver_label}"
    )

    if tag == "NEW":
        if driver == "NONE":
            evidence_text = "4) 지표 근거: 신규 항목으로 전년 기준이 부족해 YoY/dlog 신뢰도는 낮음. 현재는 원인 미확정 상태로 추세 관찰이 우선."
        else:
            evidence_text = (
                f"4) 지표 근거: 신규 항목으로 전년 기준이 부족해 YoY/dlog 신뢰도는 낮음. "
                f"현재 {driver_label} {_fmt_driver_metric(driver, product_driver_curr)}, "
                f"BU {driver_label} {_fmt_driver_metric(driver, division_driver_curr)} 기준으로 관찰."
            )
        interpretation_text = "5) 해석: 신규 항목은 원인 확정보다 초기 안정화(클릭·전환·매출 추세) 모니터링이 우선."
    elif tag == "LOW_VOL":
        if driver == "NONE":
            evidence_text = "4) 지표 근거: 저볼륨 구간으로 YoY 변동성이 크고 유효 표본이 부족해 드라이버를 확정하기 어렵습니다."
        else:
            evidence_text = (
                f"4) 지표 근거: 저볼륨 구간으로 YoY 변동성이 큼. "
                f"{driver_label} {_fmt_driver_metric(driver, product_driver_curr)} vs BU {driver_label} {_fmt_driver_metric(driver, division_driver_curr)} 참고."
            )
        interpretation_text = "5) 해석: 저볼륨 구간은 원인 단정 대신 데이터 누적 후 재판단이 적절."
    elif tag == "GONE":
        evidence_text = (
            "4) 지표 근거: 전년 매출은 있었지만 현재 매출이 0으로 관측되어 집행 중단, 커버리지 상실, 또는 리포트 누락 가능성이 큽니다."
        )
        interpretation_text = (
            "5) 해석: 원인 지표보다 운영 상태 점검이 우선입니다. 대체 캠페인 커버리지와 리마케팅 공백을 먼저 확인해야 합니다."
        )
    elif tag == "UNDEFINED":
        evidence_text = "4) 지표 근거: 핵심 지표 계산에 필요한 분모/로그 조건이 부족해 드라이버를 확정할 수 없습니다."
        interpretation_text = "5) 해석: 원인 판단 전에 전환 액션 매핑, 집계 기간, 값 이상치 여부를 먼저 점검해야 합니다."
    else:
        if driver == "NONE":
            evidence_text = "4) 지표 근거: 유효한 dlog 신호가 부족해 핵심 지표를 확정할 수 없습니다."
            interpretation_text = "5) 해석: 지표 신호가 부족해 원인 단정보다 추세 관찰이 우선."
        else:
            evidence_text = (
                f"4) 지표 근거: {driver_label} {_fmt_driver_metric(driver, product_driver_curr)}"
                f"({_fmt_pct_yoy(driver_yoy)}), BU {driver_label} {_fmt_driver_metric(driver, division_driver_curr)}"
                f"({_fmt_pct_yoy(division_driver_yoy)}), YoY 격차 {_fmt_pp(driver_gap_pp)}. "
                f"|dlog| {driver_label} {driver_dlog_abs:.3f} (CVR {dlog_cvr_abs:.3f}, AOV {dlog_aov_abs:.3f}, CPC {dlog_cpc_abs:.3f})"
            )

            if signal_max_abs < 0.05:
                interpretation_text = "5) 해석: 지표 신호 강도가 약해 원인 단정보다 추세 관찰이 우선."
            elif driver == "CPC":
                if mode_tag == "ISSUE":
                    interpretation_text = (
                        "5) 해석: 유입 효율 이슈 신호. 보조 확인은 CTR(IMP/CLICK), 클릭 볼륨, 유입단가 추세."
                    )
                else:
                    interpretation_text = (
                        "5) 해석: 유입 효율 개선 신호. 보조 확인은 CTR(IMP/CLICK) 유지, 클릭 볼륨 유지 여부."
                    )
            elif driver == "CVR":
                if mode_tag == "ISSUE":
                    interpretation_text = (
                        "5) 해석: 전환 효율 이슈 신호. 유입 트래픽 품질(키워드/오디언스/랜딩 정합성) 점검 필요."
                    )
                else:
                    interpretation_text = (
                        "5) 해석: 전환 효율 개선 신호. 유입 트래픽 품질 개선 및 랜딩/오퍼 적합도 개선 가능성."
                    )
            else:
                if mode_tag == "ISSUE":
                    interpretation_text = (
                        "5) 해석: AOV 이슈는 유입 트래픽 믹스(저의도/저가 상품 유입 비중) 문제 신호."
                    )
                else:
                    interpretation_text = (
                        "5) 해석: AOV 개선은 고의도/고가 상품 유입 비중 개선 신호."
                    )

    detail = "\n".join([channel_text, division_text, product_text, evidence_text, interpretation_text])
    return headline, detail

def _to_report_rows(
    rows: List[Dict[str, Any]],
    df: pl.DataFrame,
    scope_maps: ScopeMaps | None = None,
) -> List[Dict[str, Any]]:
    group_map: Dict[tuple[Any, Any, Any, Any], List[Dict[str, Any]]] = {}
    for item in rows:
        gkey = (
            item.get("diagnosis_subsidiary"),
            item.get("CHANNEL"),
            item.get("DIVISION"),
            str(item.get("mode_tag", "")).upper(),
        )
        group_map.setdefault(gkey, []).append(item)

    rank_map: Dict[tuple[Any, Any, Any, Any, Any], Dict[str, int]] = {}

    def _row_roas_delta(item: Dict[str, Any]) -> float:
        roas_curr = item.get("ROAS_curr")
        roas_prev = item.get("ROAS_prev")
        if roas_curr is None or roas_prev is None:
            rev_curr = _to_float(item.get("Revenue_curr_sum"))
            rev_prev = _to_float(item.get("Revenue_prev_sum"))
            spend_curr = _to_float(item.get("Spend_curr_sum"))
            spend_prev = _to_float(item.get("Spend_prev_sum"))
            roas_curr = _safe_ratio(rev_curr, spend_curr)
            roas_prev = _safe_ratio(rev_prev, spend_prev)
        if roas_curr is None or roas_prev is None:
            return 0.0
        return _to_float(roas_curr) - _to_float(roas_prev)

    for gkey, items in group_map.items():
        mode = gkey[3]
        if mode == "ISSUE":
            ordered = sorted(items, key=lambda r: (_row_roas_delta(r), str(r.get("PRODUCT", ""))))
        else:
            ordered = sorted(items, key=lambda r: (-_row_roas_delta(r), str(r.get("PRODUCT", ""))))
        pool_size = len(ordered)
        for idx, it in enumerate(ordered, start=1):
            rkey = (
                it.get("diagnosis_subsidiary"),
                it.get("CHANNEL"),
                it.get("DIVISION"),
                it.get("PRODUCT"),
                str(it.get("mode_tag", "")).upper(),
            )
            rank_map[rkey] = {"rank": idx, "pool_size": pool_size}

    channel_map, division_map, subsidiary_map = scope_maps if scope_maps is not None else _scope_maps(df)
    report_rows: List[Dict[str, Any]] = []
    for row in rows:
        channel_key = tuple(row.get(dim) for dim in CHANNEL_DIMS)
        division_key = tuple(row.get(dim) for dim in DIVISION_DIMS)
        channel_row = channel_map.get(channel_key, {})
        division_row = division_map.get(division_key, {})
        impact_adj = _to_float(row.get("Impact_adj"))
        mode_tag = str(row.get("mode_tag", "")).upper()
        topic_roas_curr = row.get("ROAS_curr")
        topic_roas_prev = row.get("ROAS_prev")
        if topic_roas_curr is None:
            topic_roas_curr = _safe_ratio(_to_float(row.get("Revenue_curr_sum")), _to_float(row.get("Spend_curr_sum")))
        if topic_roas_prev is None:
            topic_roas_prev = _safe_ratio(_to_float(row.get("Revenue_prev_sum")), _to_float(row.get("Spend_prev_sum")))
        topic_roas_delta = (
            None
            if topic_roas_curr is None or topic_roas_prev is None
            else _to_float(topic_roas_curr) - _to_float(topic_roas_prev)
        )

        division_roas_curr = _safe_ratio(
            _to_float(division_row.get("Revenue_curr_sum")),
            _to_float(division_row.get("Spend_curr_sum")),
        )
        division_roas_prev = _safe_ratio(
            _to_float(division_row.get("Revenue_prev_sum")),
            _to_float(division_row.get("Spend_prev_sum")),
        )
        division_roas_delta = (
            None
            if division_roas_curr is None or division_roas_prev is None
            else _to_float(division_roas_curr) - _to_float(division_roas_prev)
        )

        impact_contribution_pct: float | None = None
        if (
            topic_roas_delta is not None
            and division_roas_delta is not None
            and abs(division_roas_delta) > 0
        ):
            denominator = abs(division_roas_delta)
            if mode_tag == "ISSUE":
                impact_contribution_pct = max(0.0, -topic_roas_delta) / denominator
            elif mode_tag in {"IMPROVE", "DEFENSE"}:
                impact_contribution_pct = max(0.0, topic_roas_delta) / denominator
            else:
                impact_contribution_pct = abs(topic_roas_delta) / denominator
        else:
            division_delta = (
                0.0
                if division_roas_delta is None
                else division_roas_delta
            )
            impact_contribution_pct = _direct_contribution_pct(impact_adj, division_delta, mode_tag)
        summary_text, detail_text = _insight_text(
            row,
            channel_map=channel_map,
            division_map=division_map,
            subsidiary_map=subsidiary_map,
            impact_contribution_pct=impact_contribution_pct,
        )
        final_node = _final_drill_node(row.get("drill_trace"))
        rank_info = rank_map.get(
            (
                row.get("diagnosis_subsidiary"),
                row.get("CHANNEL"),
                row.get("DIVISION"),
                row.get("PRODUCT"),
                mode_tag,
            ),
            {"rank": 1, "pool_size": 1},
        )
        if mode_tag == "ISSUE":
            contribution_direction_kr = "하락"
        else:
            contribution_direction_kr = "개선"

        caution = ""
        tag = str(row.get("tag", "") or "")
        if tag == "NEW":
            caution = " 신규 토픽(NEW): 기간 확장 후 추세 재확인 권장."
        elif tag == "LOW_VOL":
            caution = " 저볼륨(LOW_VOL): 변동성 주의 해석 필요."

        why_selected_text = (
            f"{row.get('CHANNEL')}/{row.get('DIVISION')} ROAS {contribution_direction_kr}에 영향 "
            f"{rank_info['rank']}위, {_fmt_pct(impact_contribution_pct, signed=False)}, "
            f"주요 지표는 {row.get('primary_driver', 'NONE')}.{caution}"
        )
        report_rows.append(
            {
                "diagnosis_subsidiary": row.get("diagnosis_subsidiary"),
                "path": row.get("diagnosis_path"),
                "final_drill_result": {
                    "level": final_node.get("level"),
                    "rule": final_node.get("rule"),
                    "selected_path": final_node.get("selected_path"),
                },
                "impact_contribution_pct": impact_contribution_pct,
                "impact_contribution_text": _fmt_pct(impact_contribution_pct, signed=False),
                "mode_tag": row.get("mode_tag"),
                "primary_driver": row.get("primary_driver"),
                "tag": row.get("tag"),
                "selection_rank": rank_info["rank"],
                "selection_pool_size": rank_info["pool_size"],
                "why_selected_text": why_selected_text,
                "campaign_diagnosis_mode": "MANUAL",
                "summary_text": summary_text,
                "detail_text": detail_text,
                "recommended_actions": row.get("recommended_actions"),
                "action_checklist": row.get("action_checklist"),
                "dlog_CVR": row.get("dlog_CVR"),
                "dlog_AOV": row.get("dlog_AOV"),
                "dlog_CPC": row.get("dlog_CPC"),
                "CVR_curr": row.get("CVR_curr"),
                "CVR_prev": row.get("CVR_prev"),
                "CPC_curr": row.get("CPC_curr"),
                "CPC_prev": row.get("CPC_prev"),
                "AOV_curr": row.get("AOV_curr"),
                "AOV_prev": row.get("AOV_prev"),
                "topic_roas_curr": topic_roas_curr,
                "topic_roas_prev": topic_roas_prev,
                "topic_roas_delta": topic_roas_delta,
            }
        )
    return report_rows


def _top3_by_subsidiary(rows: List[Dict[str, Any]], descending: bool) -> Dict[str, List[Dict[str, Any]]]:
    return reporting_selectors.top3_by_subsidiary(
        rows,
        descending=descending,
        min_topic_roas_abs_delta=MIN_TOPIC_ROAS_ABS_DELTA,
    )


def _subsidiary_perf_map(df: pl.DataFrame) -> Dict[str, Dict[str, Any]]:
    rows = df.group_by(["SUBSIDIARY"]).agg(_sum_aggregations_full()).to_dicts()
    return {str(row.get("SUBSIDIARY", "")): row for row in rows}


def _subsidiary_metrics(scope_row: Dict[str, Any]) -> Dict[str, Any]:
    rev_curr = _to_float(scope_row.get("Revenue_curr_sum"))
    rev_prev = _to_float(scope_row.get("Revenue_prev_sum"))
    spend_curr = _to_float(scope_row.get("Spend_curr_sum"))
    spend_prev = _to_float(scope_row.get("Spend_prev_sum"))
    clicks_curr = _to_float(scope_row.get("Clicks_curr_sum"))
    clicks_prev = _to_float(scope_row.get("Clicks_prev_sum"))
    orders_curr = _to_float(scope_row.get("Orders_curr_sum"))
    orders_prev = _to_float(scope_row.get("Orders_prev_sum"))

    cpc_curr = _safe_ratio(spend_curr, clicks_curr)
    cpc_prev = _safe_ratio(spend_prev, clicks_prev)
    cvr_curr = _safe_ratio(orders_curr, clicks_curr)
    cvr_prev = _safe_ratio(orders_prev, clicks_prev)
    roas_curr = _safe_ratio(rev_curr, spend_curr)
    roas_prev = _safe_ratio(rev_prev, spend_prev)

    return {
        "rev_curr": rev_curr,
        "rev_prev": rev_prev,
        "rev_delta": rev_curr - rev_prev,
        "rev_yoy": _safe_pct_change(rev_curr, rev_prev),
        "spend_curr": spend_curr,
        "spend_prev": spend_prev,
        "spend_yoy": _safe_pct_change(spend_curr, spend_prev),
        "cpc_curr": cpc_curr,
        "cpc_prev": cpc_prev,
        "cpc_yoy": _safe_pct_change(cpc_curr, cpc_prev),
        "cvr_curr": cvr_curr,
        "cvr_prev": cvr_prev,
        "cvr_yoy": _safe_pct_change(cvr_curr, cvr_prev),
        "roas_curr": roas_curr,
        "roas_prev": roas_prev,
        "roas_delta": None if roas_curr is None or roas_prev is None else roas_curr - roas_prev,
        "roas_yoy": _safe_pct_change(roas_curr, roas_prev),
    }


def _choose_subsidiary_mode(
    subsidiary: str,
    perf: Dict[str, Any],
    issues_by_sub: Dict[str, List[Dict[str, Any]]],
    improves_by_sub: Dict[str, List[Dict[str, Any]]],
) -> tuple[str, List[Dict[str, Any]]]:
    return reporting_selectors.choose_subsidiary_mode(subsidiary, perf, issues_by_sub, improves_by_sub)


def _subsidiary_overall_comment(subsidiary: str, perf: Dict[str, Any], mode: str) -> str:
    return reporting_rendering.subsidiary_overall_comment(subsidiary, perf, mode)


def _subsidiary_detail_comment(mode: str, top3: List[Dict[str, Any]]) -> str:
    return reporting_rendering.subsidiary_detail_comment(mode, top3)


def _build_subsidiary_reports(
    df: pl.DataFrame,
    subsidiaries: List[str],
    issues_top3_by_sub: Dict[str, List[Dict[str, Any]]],
    improvements_top3_by_sub: Dict[str, List[Dict[str, Any]]],
) -> List[Dict[str, Any]]:
    perf_map = _subsidiary_perf_map(df)
    reports: List[Dict[str, Any]] = []

    for subsidiary in subsidiaries:
        sub = str(subsidiary)
        perf = _subsidiary_metrics(perf_map.get(sub, {}))
        mode, selected_top3 = _choose_subsidiary_mode(
            sub,
            perf,
            issues_top3_by_sub,
            improvements_top3_by_sub,
        )
        reports.append(
            {
                "SUBSIDIARY": sub,
                "mode": mode,
                "overall_comment": _subsidiary_overall_comment(sub, perf, mode),
                "detailed_comment": _subsidiary_detail_comment(mode, selected_top3),
                "performance": perf,
                "top3": selected_top3,
            }
    )
    return reports


def _topic_scope(row: Dict[str, Any]) -> tuple[str, str, str]:
    path = row.get("path", {})
    if not isinstance(path, dict):
        path = {}
    channel = str(path.get("CHANNEL", "") or "")
    division = str(path.get("DIVISION", "") or "")
    product = str(path.get("PRODUCT", "") or "")
    return channel, division, product


def _topic_root_cause_text(row: Dict[str, Any]) -> str:
    driver = str(row.get("primary_driver", "NONE") or "NONE").upper()
    tag = str(row.get("tag", "") or "").upper()

    if tag == "NEW":
        return "신규 토픽(NEW): 전년 기준이 부족해 원인 확정보다 초기 추세 모니터링이 우선."
    if tag == "LOW_VOL":
        return "저볼륨(LOW_VOL): 표본이 작아 변동성이 커 드라이버 해석 신뢰도에 주의 필요."
    if tag == "GONE":
        return "종료(GONE): 전년 매출은 있으나 현재 매출 0으로 집행중단/커버리지 공백 여부 점검 필요."
    if tag == "UNDEFINED":
        return "미정의(UNDEFINED): 분모/로그 조건 부족으로 드라이버 확정이 어려워 데이터 품질 점검 우선."

    if driver == "CVR":
        curr_key, prev_key, dlog_key = "CVR_curr", "CVR_prev", "dlog_CVR"
    elif driver == "AOV":
        curr_key, prev_key, dlog_key = "AOV_curr", "AOV_prev", "dlog_AOV"
    elif driver == "CPC":
        curr_key, prev_key, dlog_key = "CPC_curr", "CPC_prev", "dlog_CPC"
    else:
        return "핵심 지표 미확정(NONE): dlog 신호가 약해 원인 단정보다 추세 관찰이 우선."

    curr = row.get(curr_key)
    prev = row.get(prev_key)
    yoy = _safe_pct_change(curr, prev)
    dlog_raw = row.get(dlog_key)
    dlog_text = "N/A" if dlog_raw is None else f"{_to_float(dlog_raw):.3f}"

    why_selected = str(row.get("why_selected_text", "") or "").strip()
    detail = f"{driver} 기반 원인. YoY {_fmt_pct_yoy(yoy)}, dlog {dlog_text}."
    if why_selected:
        detail = f"{detail} 선정근거: {why_selected}"
    return detail


def _subsidiary_issue_action_sheet_df(
    subsidiary: str,
    rows: List[Dict[str, Any]],
    source_label: str,
) -> pl.DataFrame:
    columns = [
        "Subsidiary",
        "Rank",
        "Media",
        "BU",
        "Product",
        "Issue Finding",
        "Root Cause Analysis",
        "Action Item",
        "Checklist",
        "Primary Driver",
        "Tag",
        "Contribution",
        "Source",
    ]
    if not rows:
        return pl.DataFrame({col: [] for col in columns})

    out: List[Dict[str, Any]] = []
    for idx, row in enumerate(rows, start=1):
        channel, division, product = _topic_scope(row)
        issue_finding = str(row.get("summary_text", "") or row.get("why_selected_text", "") or "").strip()
        root_cause = _topic_root_cause_text(row)
        action_item = str(row.get("recommended_actions", "") or "").strip()
        checklist = str(row.get("action_checklist", "") or "").strip()
        out.append(
            {
                "Subsidiary": subsidiary,
                "Rank": idx,
                "Media": channel,
                "BU": division,
                "Product": product,
                "Issue Finding": issue_finding,
                "Root Cause Analysis": root_cause,
                "Action Item": action_item,
                "Checklist": checklist,
                "Primary Driver": str(row.get("primary_driver", "NONE") or "NONE"),
                "Tag": str(row.get("tag", "") or ""),
                "Contribution": str(row.get("impact_contribution_text", "N/A") or "N/A"),
                "Source": source_label,
            }
        )
    return pl.DataFrame(out).select(columns)


def _build_subsidiary_issue_action_sheets(
    subsidiaries: List[str],
    subsidiary_reports: List[Dict[str, Any]],
    issues_top3_by_sub: Dict[str, List[Dict[str, Any]]],
    improvements_top3_by_sub: Dict[str, List[Dict[str, Any]]],
) -> Dict[str, pl.DataFrame]:
    report_map: Dict[str, Dict[str, Any]] = {}
    for report in subsidiary_reports:
        if not isinstance(report, dict):
            continue
        sub = str(report.get("SUBSIDIARY", "") or "")
        if sub:
            report_map[sub] = report

    sheet_map: Dict[str, pl.DataFrame] = {}
    for sub in subsidiaries:
        issue_rows = issues_top3_by_sub.get(sub, [])
        if issue_rows:
            selected_rows = issue_rows
            source_label = "issues_top3_products_by_subsidiary"
        else:
            report_top3 = []
            report = report_map.get(sub, {})
            if isinstance(report, dict):
                maybe_top3 = report.get("top3", [])
                if isinstance(maybe_top3, list):
                    report_top3 = [row for row in maybe_top3 if isinstance(row, dict)]
            if report_top3:
                selected_rows = report_top3
                source_label = "subsidiary_reports.top3"
            else:
                improve_rows = improvements_top3_by_sub.get(sub, [])
                selected_rows = improve_rows
                source_label = "improvements_top3_products_by_subsidiary"

        # Excel sheet names are max 31 chars.
        sheet_name = str(sub)[:31] if sub else "UNKNOWN"
        sheet_map[sheet_name] = _subsidiary_issue_action_sheet_df(sub, selected_rows, source_label)
    return sheet_map


def _template_topics_for_sub(
    subsidiary: str,
    issues_top3_by_sub: Dict[str, List[Dict[str, Any]]],
    improvements_top3_by_sub: Dict[str, List[Dict[str, Any]]],
    report_map: Dict[str, Dict[str, Any]],
    preferred_mode: str,
) -> List[Dict[str, Any]]:
    mode = str(preferred_mode or "ISSUE").upper()
    issue_rows = issues_top3_by_sub.get(subsidiary, [])[:3]
    improve_rows = improvements_top3_by_sub.get(subsidiary, [])[:3]
    if mode == "IMPROVE":
        if improve_rows:
            return improve_rows
    else:
        if issue_rows:
            return issue_rows

    report = report_map.get(subsidiary, {})
    if isinstance(report, dict):
        top3 = report.get("top3", [])
        if isinstance(top3, list) and top3:
            return [row for row in top3 if isinstance(row, dict)][:3]

    return issue_rows if issue_rows else improve_rows


def _template_issue_text(row: Dict[str, Any]) -> str:
    channel, division, product = _topic_scope(row)
    summary = str(row.get("summary_text", "") or row.get("why_selected_text", "") or "").strip()
    root_cause = _topic_root_cause_text(row)
    prefix = "/".join([token for token in [channel, division, product] if token])
    parts = []
    if prefix:
        parts.append(prefix)
    if summary:
        parts.append(summary)
    if root_cause:
        parts.append(f"원인분석: {root_cause}")
    return " | ".join(parts)


def _template_action_text(row: Dict[str, Any]) -> str:
    action = str(row.get("recommended_actions", "") or "").strip()
    checklist = str(row.get("action_checklist", "") or "").strip()
    if action and checklist:
        return f"{action} | {checklist}"
    if action:
        return action
    if checklist:
        return checklist
    return "액션 없음"


def _save_reff_format_workbook(
    project_root: Path,
    subsidiaries: List[str],
    subsidiary_reports: List[Dict[str, Any]],
    issues_top3_by_sub: Dict[str, List[Dict[str, Any]]],
    improvements_top3_by_sub: Dict[str, List[Dict[str, Any]]],
    comparison_meta: Dict[str, Any],
) -> tuple[bool, str, Path]:
    template_path = project_root.parent / "Reff" / "ROAS Report Format.xlsx"
    output_path = project_root / "output" / "ROAS Report Format_filled.xlsx"
    if not template_path.exists():
        return False, f"Template not found: {template_path}", output_path

    try:
        from openpyxl import load_workbook
    except Exception as exc:  # pragma: no cover
        return False, f"openpyxl import failed: {exc}", output_path

    report_map: Dict[str, Dict[str, Any]] = {}
    for report in subsidiary_reports:
        if not isinstance(report, dict):
            continue
        sub = str(report.get("SUBSIDIARY", "") or "")
        if sub:
            report_map[sub] = report

    curr_year = int(comparison_meta.get("curr_year", date.today().year))
    prev_year = int(comparison_meta.get("prev_year", curr_year - 1))
    mtd_month = comparison_meta.get("mtd_month_cutoff") or comparison_meta.get("mtd_month_start")
    mtd_day = comparison_meta.get("mtd_day_cutoff")
    if mtd_month and mtd_day:
        period_text = f"({curr_year} MTD 1/1 - {int(mtd_month)}/{int(mtd_day)})"
    elif mtd_month:
        period_text = f"({curr_year} MTD month {int(mtd_month)})"
    else:
        period_text = f"({curr_year} vs {prev_year})"
    source_line = f"* Source: GMPD RAW, LMDI pipeline, {prev_year} & {curr_year} {period_text}"
    method_line = "** Data Source: Impact(LMDI) + RCA(CVR/AOV/CPC) + rule-based action recommendation"

    try:
        workbook = load_workbook(template_path)
        for sub in subsidiaries:
            if sub not in workbook.sheetnames:
                continue
            worksheet = workbook[sub]
            report = report_map.get(sub, {})
            mode = str(report.get("mode", "ISSUE") if isinstance(report, dict) else "ISSUE").upper()
            rows = _template_topics_for_sub(
                sub,
                issues_top3_by_sub=issues_top3_by_sub,
                improvements_top3_by_sub=improvements_top3_by_sub,
                report_map=report_map,
                preferred_mode=mode,
            )

            worksheet.cell(row=2, column=2).value = source_line
            worksheet.cell(row=3, column=2).value = method_line

            issue_texts = [_template_issue_text(row) for row in rows[:3]]
            action_texts = [_template_action_text(row) for row in rows[:3]]

            for idx in range(3):
                issue_row = 8 + idx
                action_row = 15 + idx
                worksheet.cell(row=issue_row, column=3).value = issue_texts[idx] if idx < len(issue_texts) else None
                worksheet.cell(row=action_row, column=3).value = action_texts[idx] if idx < len(action_texts) else None

        if "Sheet1" in workbook.sheetnames:
            sheet1 = workbook["Sheet1"]
            for row_idx in range(3, sheet1.max_row + 1):
                sub_name = str(sheet1.cell(row=row_idx, column=2).value or "").strip()
                if not sub_name:
                    continue
                if sub_name not in workbook.sheetnames:
                    continue
                report = report_map.get(sub_name, {})
                mode = str(report.get("mode", "ISSUE") if isinstance(report, dict) else "ISSUE").upper()
                rows = _template_topics_for_sub(
                    sub_name,
                    issues_top3_by_sub=issues_top3_by_sub,
                    improvements_top3_by_sub=improvements_top3_by_sub,
                    report_map=report_map,
                    preferred_mode=mode,
                )
                mode_text = "개선" if mode == "IMPROVE" else "이슈"
                sheet1.cell(row=row_idx, column=4).value = mode_text
                sheet1.cell(row=row_idx + 1, column=4).value = "원인 분석" if mode == "ISSUE" else "개선 분석"
                sheet1.cell(row=row_idx + 2, column=4).value = "Action Items"

                summary_lines: List[str] = []
                cause_lines: List[str] = []
                action_lines: List[str] = []
                for idx, topic in enumerate(rows[:3], start=1):
                    summary_text = str(topic.get("summary_text", "") or topic.get("why_selected_text", "") or "").strip()
                    if not summary_text:
                        summary_text = _template_issue_text(topic)
                    channel, division, product = _topic_scope(topic)
                    driver = str(topic.get("primary_driver", "NONE") or "NONE")
                    summary_lines.append(f"{idx}. {summary_text}")
                    cause_lines.append(
                        f"{idx}. {channel}/{division}/{product} | driver={driver} | {_topic_root_cause_text(topic)}"
                    )
                    action_lines.append(f"{idx}. {_template_action_text(topic)}")

                sheet1.cell(row=row_idx, column=5).value = "\n".join(summary_lines) if summary_lines else "토픽 없음"
                sheet1.cell(row=row_idx + 1, column=5).value = "\n".join(cause_lines) if cause_lines else "원인 없음"
                sheet1.cell(row=row_idx + 2, column=5).value = "\n".join(action_lines) if action_lines else "액션 없음"

        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output_path)
        workbook.close()
        return True, "", output_path
    except PermissionError as exc:
        return False, str(exc), output_path
    except Exception as exc:
        return False, str(exc), output_path


def _campaign_sheet_df(rows: List[Dict[str, Any]]) -> pl.DataFrame:
    columns = [
        "diagnosis_subsidiary",
        *DIMENSIONS,
        "Revenue_curr_sum",
        "Revenue_prev_sum",
        "Spend_curr_sum",
        "Spend_prev_sum",
        "Clicks_curr_sum",
        "Clicks_prev_sum",
        "Orders_curr_sum",
        "Orders_prev_sum",
        "Impact_raw",
        "Impact_adj",
        "Share_childsum",
        "Share_parentdelta",
        "Residual_parent",
        "Residual_applied_flag",
        "mode_tag",
        "primary_driver",
        "tag",
        "dlog_CVR",
        "dlog_AOV",
        "dlog_CPC",
        "total_dlog_roas",
        "CPC_curr",
        "CPC_prev",
        "CVR_curr",
        "CVR_prev",
        "AOV_curr",
        "AOV_prev",
        "ROAS_curr",
        "ROAS_prev",
        "campaign_diagnosis_mode",
        "recommended_actions",
        "action_checklist",
        "drill_trace_text",
    ]
    if not rows:
        return pl.DataFrame({col: [] for col in columns})

    normalized_rows: list[dict[str, Any]] = []
    for row in rows:
        normalized_rows.append({col: row.get(col) for col in columns})
    return pl.DataFrame(normalized_rows).select(columns)


def _campaign_pretty_sheet_df(rows: List[Dict[str, Any]], curr_year: int, prev_year: int) -> pl.DataFrame:
    revenue_curr_col = f"revenue_{curr_year}"
    revenue_prev_col = f"revenue_{prev_year}"
    spend_curr_col = f"spend_{curr_year}"
    spend_prev_col = f"spend_{prev_year}"
    roas_curr_col = f"roas_{curr_year}"
    roas_prev_col = f"roas_{prev_year}"
    columns = [
        "subsidiary",
        "media",
        "bu",
        "product",
        "mode",
        "driver",
        "impact_roas",
        revenue_curr_col,
        revenue_prev_col,
        "revenue_yoy",
        spend_curr_col,
        spend_prev_col,
        "spend_yoy",
        roas_curr_col,
        roas_prev_col,
        "roas_yoy",
        "tag",
        "action_one",
        "checklist",
        "trace",
    ]
    if not rows:
        return pl.DataFrame({col: [] for col in columns})

    pretty_rows: List[Dict[str, Any]] = []
    for row in rows:
        rev_curr = _to_float(row.get("Revenue_curr_sum"))
        rev_prev = _to_float(row.get("Revenue_prev_sum"))
        spend_curr = _to_float(row.get("Spend_curr_sum"))
        spend_prev = _to_float(row.get("Spend_prev_sum"))
        roas_curr = row.get("ROAS_curr")
        roas_prev = row.get("ROAS_prev")
        if roas_curr is None:
            roas_curr = _safe_ratio(rev_curr, spend_curr)
        if roas_prev is None:
            roas_prev = _safe_ratio(rev_prev, spend_prev)

        pretty_rows.append(
            {
                "subsidiary": row.get("diagnosis_subsidiary"),
                "media": row.get("CHANNEL"),
                "bu": row.get("DIVISION"),
                "product": row.get("PRODUCT"),
                "mode": row.get("mode_tag"),
                "driver": row.get("primary_driver"),
                "impact_roas": (
                    "N/A"
                    if row.get("Impact_adj") is None
                    else f"{_to_float(row.get('Impact_adj')):+.3f}"
                ),
                revenue_curr_col: _fmt_money(rev_curr),
                revenue_prev_col: _fmt_money(rev_prev),
                "revenue_yoy": _fmt_pct(_safe_pct_change(rev_curr, rev_prev)),
                spend_curr_col: _fmt_money(spend_curr),
                spend_prev_col: _fmt_money(spend_prev),
                "spend_yoy": _fmt_pct(_safe_pct_change(spend_curr, spend_prev)),
                roas_curr_col: _fmt_roas(roas_curr),
                roas_prev_col: _fmt_roas(roas_prev),
                "roas_yoy": _fmt_pct(_safe_pct_change(roas_curr, roas_prev)),
                "tag": row.get("tag"),
                "action_one": row.get("recommended_actions"),
                "checklist": row.get("action_checklist"),
                "trace": row.get("drill_trace_text"),
            }
        )

    return pl.DataFrame(pretty_rows).select(columns)


def _division_sheet_df(divisions: List[Dict[str, Any]]) -> pl.DataFrame:
    columns = [
        "SUBSIDIARY",
        "DIVISION",
        "Revenue_curr_sum",
        "Revenue_prev_sum",
        "Spend_curr_sum",
        "Spend_prev_sum",
        "Revenue_delta",
        "Revenue_yoy_pct",
        "Spend_yoy_pct",
        "ROAS_curr",
        "ROAS_prev",
        "ROAS_delta",
        "ROAS_yoy_pct",
        "Impact_raw",
        "Impact_adj",
        "Impact_raw_method",
        "Residual_parent",
        "Residual_applied_flag",
        "Parent_roas_delta",
        "Contribution_to_total_delta_pct",
        "Contribution_to_decline_pct",
        "Contribution_to_growth_pct",
        "Contribution_label",
        "summary_text",
    ]
    flat_rows: List[Dict[str, Any]] = []
    for block in divisions:
        if not isinstance(block, dict):
            continue
        rows = block.get("rows", [])
        if not isinstance(rows, list):
            continue
        for row in rows:
            if not isinstance(row, dict):
                continue
            flat_rows.append({col: row.get(col) for col in columns})
    if not flat_rows:
        return pl.DataFrame({col: [] for col in columns})
    return pl.DataFrame(flat_rows).select(columns)


def _flatten_trace(drill_trace: Dict[str, Any]) -> pl.DataFrame:
    columns = [
        "step",
        "focus_SUBSIDIARY",
        "level",
        "rule",
        *[f"parent_{dim}" for dim in DIMENSIONS],
        "parent_delta",
        "top1_share_childsum",
        "top2_share_childsum",
        "impact_floor",
        "selected_paths",
    ]
    events = drill_trace.get("events", []) if isinstance(drill_trace, dict) else []
    if not isinstance(events, list) or not events:
        return pl.DataFrame({col: [] for col in columns})

    rows: list[dict[str, Any]] = []
    for idx, event in enumerate(events, start=1):
        if not isinstance(event, dict):
            continue
        parent_path = event.get("parent_path", {})
        if not isinstance(parent_path, dict):
            parent_path = {}
        selected_paths = event.get("selected_paths", [])
        selected_tokens: list[str] = []
        if isinstance(selected_paths, list):
            for path_obj in selected_paths:
                if isinstance(path_obj, dict):
                    token = "/".join(
                        str(path_obj.get(dim, "")) for dim in DIMENSIONS if path_obj.get(dim) is not None
                    )
                    selected_tokens.append(token)
        row_data = {
            "step": idx,
            "focus_SUBSIDIARY": event.get("focus_subsidiary"),
            "level": event.get("level"),
            "rule": event.get("rule"),
            "parent_delta": event.get("parent_delta"),
            "top1_share_childsum": event.get("top1_share_childsum"),
            "top2_share_childsum": event.get("top2_share_childsum"),
            "impact_floor": event.get("impact_floor"),
            "selected_paths": " ; ".join(selected_tokens),
        }
        for dim in DIMENSIONS:
            row_data[f"parent_{dim}"] = parent_path.get(dim)
        rows.append(row_data)

    if not rows:
        return pl.DataFrame({col: [] for col in columns})
    return pl.DataFrame(rows).select(columns)


def run_reporting_pipeline(
    input_path: str | Path,
    curr_year: int | None = None,
    prev_year: int | None = None,
) -> None:
    pipeline_start = perf_counter()
    stage_start = pipeline_start
    stage_timings: list[tuple[str, float]] = []
    current_year, previous_year = _resolve_comparison_years(curr_year, prev_year)

    def _mark(stage_name: str) -> None:
        nonlocal stage_start
        now = perf_counter()
        stage_timings.append((stage_name, now - stage_start))
        stage_start = now

    project_root = Path(__file__).resolve().parents[2]
    input_path = Path(input_path).expanduser().resolve()
    output_json_path = project_root / "output" / "summary.json"
    output_excel_path = project_root / "output" / "summary.xlsx"
    output_raw_calc_excel_path = project_root / "output" / "raw_html_mtd_conversion_extrevenue.xlsx"
    output_html_path = project_root / "output" / "summary.html"

    df, comparison_meta = load_input_frame(
        input_path,
        curr_year=current_year,
        prev_year=previous_year,
        cache_root=project_root,
    )
    _mark("load_input_frame")
    analysis = run_subsidiary_analysis(df, dimensions=DIMENSIONS)
    _mark("run_subsidiary_analysis")
    subsidiaries = analysis.subsidiaries
    issues_with_actions = analysis.issues_with_actions
    improvements_with_actions = analysis.improvements_with_actions
    trace_by_subsidiary = analysis.trace_by_subsidiary

    scope_maps = _scope_maps(df)
    issues_report_rows = _to_report_rows(issues_with_actions, df=df, scope_maps=scope_maps)
    improvements_report_rows = _to_report_rows(improvements_with_actions, df=df, scope_maps=scope_maps)
    issues_top3_by_sub = _top3_by_subsidiary(issues_report_rows, descending=True)
    improvements_top3_by_sub = _top3_by_subsidiary(improvements_report_rows, descending=True)
    division_parallel = _division_parallel_contribution(df)
    subsidiary_reports = _build_subsidiary_reports(
        df,
        subsidiaries=[str(sub) for sub in subsidiaries],
        issues_top3_by_sub=issues_top3_by_sub,
        improvements_top3_by_sub=improvements_top3_by_sub,
    )
    subsidiary_issue_action_sheets = _build_subsidiary_issue_action_sheets(
        subsidiaries=[str(sub) for sub in subsidiaries],
        subsidiary_reports=subsidiary_reports,
        issues_top3_by_sub=issues_top3_by_sub,
        improvements_top3_by_sub=improvements_top3_by_sub,
    )
    _mark("build_report_rows")

    run_metadata = _build_run_metadata(input_path, current_year=current_year, comparison_meta=comparison_meta)
    summary = {
        "week_key": run_metadata["week_key"],
        "generated_at": run_metadata["generated_at"],
        "source_hash": run_metadata["source_hash"],
        "comparison_meta": comparison_meta,
        "subsidiaries_analyzed": subsidiaries,
        "decomposition_method": {
            "impact_raw": "LMDI(log-mean) with direct delta fallback when log condition is invalid",
            "impact_adj": "Impact_raw + residual correction proportionally allocated by |Impact_raw|",
            "division_contribution_pct": "Contribution_to_decline_pct/growth_pct is based on Impact_adj vs parent ROAS delta",
        },
        "division_parallel_contribution": division_parallel,
        "subsidiary_reports": subsidiary_reports,
        "issues_top3_products_by_subsidiary": issues_top3_by_sub,
        "improvements_top3_products_by_subsidiary": improvements_top3_by_sub,
        "issues_top_products": issues_report_rows,
        "improvements_top_products": improvements_report_rows,
    }
    _mark("build_summary")

    save_summary_json(output_json_path, summary)
    save_summary_html(output_html_path, summary, df)
    _mark("save_json_html")

    issues_df = _campaign_sheet_df(issues_with_actions)
    improvements_df = _campaign_sheet_df(improvements_with_actions)
    issues_view_df = _campaign_pretty_sheet_df(issues_with_actions, curr_year=current_year, prev_year=previous_year)
    improvements_view_df = _campaign_pretty_sheet_df(
        improvements_with_actions,
        curr_year=current_year,
        prev_year=previous_year,
    )
    division_df = _division_sheet_df(division_parallel)
    combined_events: list[dict[str, Any]] = []
    for subsidiary, drill_trace in trace_by_subsidiary.items():
        events = drill_trace.get("events", []) if isinstance(drill_trace, dict) else []
        if not isinstance(events, list):
            continue
        for event in events:
            if not isinstance(event, dict):
                continue
            item = dict(event)
            item["focus_subsidiary"] = subsidiary
            combined_events.append(item)
    trace_df = _flatten_trace({"focus_path": {"SUBSIDIARY": "ALL"}, "events": combined_events})
    _mark("build_excel_frames")
    excel_saved, excel_error_message = save_output_workbook(
        output_excel_path,
        {
            "issues": issues_df,
            "improvements": improvements_df,
            "issues_view": issues_view_df,
            "improvements_view": improvements_view_df,
            "division_parallel": division_df,
            "trace": trace_df,
            **subsidiary_issue_action_sheets,
        },
    )
    raw_calc_saved, raw_calc_error_message, raw_calc_meta, raw_calc_saved_path = save_html_calc_raw_workbook(
        output_raw_calc_excel_path,
        input_path=input_path,
        curr_year=current_year,
        prev_year=previous_year,
    )
    _mark("save_excel")
    archive_saved, archive_error_message, archive_dir = _archive_outputs(
        project_root,
        week_key=run_metadata["week_key"],
        output_paths=[output_json_path, output_html_path]
        + ([output_excel_path] if excel_saved else [])
        + ([raw_calc_saved_path] if raw_calc_saved else []),
    )
    reff_saved, reff_error_message, reff_output_path = _save_reff_format_workbook(
        project_root=project_root,
        subsidiaries=[str(sub) for sub in subsidiaries],
        subsidiary_reports=subsidiary_reports,
        issues_top3_by_sub=issues_top3_by_sub,
        improvements_top3_by_sub=improvements_top3_by_sub,
        comparison_meta=comparison_meta,
    )
    if reff_saved:
        archive_saved_reff, archive_reff_error, _ = _archive_outputs(
            project_root,
            week_key=run_metadata["week_key"],
            output_paths=[reff_output_path],
        )
        if not archive_saved_reff and not archive_error_message:
            archive_error_message = archive_reff_error
    total_elapsed = perf_counter() - pipeline_start

    print(
        "Summary prepared: "
        f"subsidiaries={len(subsidiaries)}, "
        f"issue_topics={len(issues_report_rows)}, "
        f"improve_topics={len(improvements_report_rows)}"
    )
    stage_text = ", ".join([f"{name}={seconds:.3f}s" for name, seconds in stage_timings])
    print(f"Stage Timing: {stage_text}")
    print(f"Total Elapsed: {total_elapsed:.3f}s")
    print(f"Comparison Meta: {comparison_meta}")
    print(f"Saved JSON: {output_json_path}")
    print(f"Saved HTML: {output_html_path}")
    if excel_saved:
        print(f"Saved Excel: {output_excel_path}")
    else:
        print(f"Excel save skipped (file may be open/locked): {excel_error_message}")
    if raw_calc_saved:
        print(
            "Saved Calc RAW Excel: "
            f"{raw_calc_saved_path} "
            f"(rows_full={raw_calc_meta.get('raw_rows_mtd_html_calc_full', 0)}, "
            f"rows_engine={raw_calc_meta.get('engine_rows_html_calc', 0)})"
        )
        if raw_calc_error_message:
            print(f"Calc RAW Excel note: {raw_calc_error_message}")
    else:
        print(f"Calc RAW Excel save skipped: {raw_calc_error_message}")
    if archive_saved:
        print(f"Archived outputs: {archive_dir}")
    else:
        print(f"Archive save skipped: {archive_error_message}")
    if reff_saved:
        print(f"Saved Reff format Excel: {reff_output_path}")
    else:
        print(f"Reff format save skipped: {reff_error_message}")



def _source_hash(path: Path) -> str:
    digest = hashlib.sha1()
    try:
        with path.open("rb") as handle:
            while True:
                chunk = handle.read(1024 * 1024)
                if not chunk:
                    break
                digest.update(chunk)
        return digest.hexdigest()
    except PermissionError:
        # OneDrive/Excel can keep a transient lock on the workbook.
        return _source_metadata_hash(path)


def _source_metadata_hash(path: Path) -> str:
    resolved = path.expanduser().resolve()
    try:
        stat = resolved.stat()
        fingerprint = f"{resolved}|{stat.st_size}|{stat.st_mtime_ns}|{stat.st_ctime_ns}"
    except OSError:
        fingerprint = str(resolved)
    return hashlib.sha1(fingerprint.encode("utf-8", errors="replace")).hexdigest()


def _build_week_key(current_year: int, comparison_meta: Dict[str, Any]) -> str:
    month_raw = comparison_meta.get("mtd_month_cutoff") or comparison_meta.get("mtd_month_start")
    day_raw = comparison_meta.get("mtd_day_cutoff")
    month = int(month_raw) if month_raw is not None else None
    day = int(day_raw) if day_raw is not None else None
    if month is not None and day is not None:
        return f"{current_year}-{month:02d}-{day:02d}"
    if month is not None:
        return f"{current_year}-{month:02d}"
    return str(current_year)


def _build_run_metadata(input_path: Path, current_year: int, comparison_meta: Dict[str, Any]) -> Dict[str, str]:
    return {
        "week_key": _build_week_key(current_year, comparison_meta),
        "generated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "source_hash": _source_hash(input_path),
    }


def _archive_outputs(project_root: Path, week_key: str, output_paths: List[Path]) -> tuple[bool, str, Path]:
    archive_dir = project_root / "data" / "archive" / week_key
    try:
        archive_dir.mkdir(parents=True, exist_ok=True)
        for output_path in output_paths:
            if output_path.exists():
                shutil.copy2(output_path, archive_dir / output_path.name)
    except OSError as exc:
        return False, str(exc), archive_dir
    return True, "", archive_dir

